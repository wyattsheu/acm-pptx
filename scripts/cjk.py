#!/usr/bin/env python3
"""East-Asian text support shared by build_from_outline.py and compose.py.

python-pptx writes every run as `lang="en-US"` with a Latin typeface only.
LibreOffice then renders 中文 with a fallback that overlaps or drops glyphs,
and PowerPoint picks whatever the system has, so a deck looks different on
every machine. The fix is two attributes PowerPoint itself writes for Chinese
text: `lang="zh-TW"` on the run and an `<a:ea typeface="..."/>` beside the
Latin face. `tag_deck()` walks every run in a deck and adds both wherever the
text carries CJK, so neither builder has to remember.
"""
from __future__ import annotations

import re

from pptx.oxml.ns import qn

CJK = re.compile(r"[　-〿぀-ヿ㐀-䶿一-鿿豈-﫿＀-￯]")
DEFAULT_FONT = "Microsoft JhengHei"
DEFAULT_LANG = "zh-TW"

# rPr children must keep schema order; `ea` sits after `latin`, before these
_AFTER_EA = ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def has_cjk(text: str) -> bool:
    return bool(CJK.search(text or ""))


def cjk_chars(text: str) -> int:
    return len(CJK.findall(text or ""))


def tag_run(r_el, font: str = DEFAULT_FONT, lang: str = DEFAULT_LANG) -> bool:
    """Set lang and the East-Asian face on one `a:r`; True if changed."""
    t = r_el.find(qn("a:t"))
    if t is None or not has_cjk(t.text or ""):
        return False
    rPr = r_el.find(qn("a:rPr"))
    if rPr is None:
        rPr = r_el.makeelement(qn("a:rPr"), {})
        r_el.insert(0, rPr)
    changed = False
    if rPr.get("lang") != lang:
        rPr.set("lang", lang)
        changed = True
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {"typeface": font})
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(ea)
        else:
            anchor = next((c for c in rPr if c.tag in {qn(n) for n in _AFTER_EA}), None)
            if anchor is not None:
                anchor.addprevious(ea)
            else:
                rPr.append(ea)
        changed = True
    elif ea.get("typeface") != font:
        ea.set("typeface", font)
        changed = True
    return changed


def tag_text_frame(tf, font: str = DEFAULT_FONT, lang: str = DEFAULT_LANG) -> int:
    n = 0
    for r_el in tf._txBody.iter(qn("a:r")):
        n += tag_run(r_el, font, lang)
    return n


def tag_slide(slide, font: str = DEFAULT_FONT, lang: str = DEFAULT_LANG) -> int:
    n = 0
    for shp in slide.shapes:
        if shp.has_text_frame:
            n += tag_text_frame(shp.text_frame, font, lang)
        if shp.has_table:
            for row in shp.table.rows:
                for cell in row.cells:
                    n += tag_text_frame(cell.text_frame, font, lang)
    if slide.has_notes_slide:
        n += tag_text_frame(slide.notes_slide.notes_text_frame, font, lang)
    return n


def tag_deck(prs, font: str = DEFAULT_FONT, lang: str = DEFAULT_LANG) -> int:
    """Every run in the deck that carries CJK gets lang + East-Asian face."""
    return sum(tag_slide(s, font, lang) for s in prs.slides)


def font_from_outline(outline: dict) -> str:
    return (outline.get("meta") or {}).get("cjk_font") or DEFAULT_FONT
