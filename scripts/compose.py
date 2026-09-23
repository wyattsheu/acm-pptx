#!/usr/bin/env python3
"""Second pass over a deck built by build_from_outline.py.

build_from_outline.py clones template slides and fills text. This adds the
things a paper talk actually needs: a red claim line under the title, a
figure with its caption, annotations drawn on that figure, a comparison
matrix, an equation band, an embedded video, and the bottom callout.

    python scripts/build_from_outline.py outline.json -o talk.pptx
    python scripts/compose.py outline.json talk.pptx

Regions are computed, never hand-written: say `"layout": "figure-right"`
and the body box shrinks to make room. Overlap is the failure mode every
other slide generator ships with; this is how it is avoided.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- palette
RED = RGBColor(0xC0, 0x00, 0x00)      # claim line, callout frame, annotations
GREY = RGBColor(0x59, 0x59, 0x59)     # captions
HEAD_BG = RGBColor(0x3E, 0x4A, 0x5B)  # matrix header
HEAD_FG = RGBColor(0xFF, 0xFF, 0xFF)
ROW_HI = RGBColor(0xE8, 0xF2, 0xE1)   # the "Ours" row
BAND = RGBColor(0xF2, 0xF2, 0xF2)     # equation band

# ---------------------------------------------------------------- geometry
# template: 13.333 x 7.5in, title 0.40-1.85, body 2.00-6.75, page no. at 6.95
# The title band has to stop above the subtitle: the title is bottom-anchored,
# so its box bottom is where the glyphs sit. 0.28 + 0.86 = 1.14, clear of 1.16.
TITLE = (0.92, 0.28, 11.50, 0.86)
SUBTITLE = (0.95, 1.16, 11.50, 0.44)
CONTENT_TOP_PLAIN = 1.72
CONTENT_TOP_SUB = 1.78
CONTENT_BOTTOM = 6.80
CALLOUT_H = 0.78
CAPTION_H = 0.34
GUTTER = 0.30
LEFT, RIGHT = 0.92, 12.42          # body text column
BLEED_L, BLEED_R = 0.62, 12.72     # figures may run wider than text
STAGE_H = 1.60                     # the recurring pipeline strip, pinned top-right
ASSERT = (0.92, 1.16, 11.50, 0.88)  # assertion-evidence: the sentence headline
ASSERT_TOP = 2.10                   # evidence starts below it
DIVIDER = RGBColor(0xC8, 0xCC, 0xD2)


def regions(spec: dict) -> dict:
    """Body box and figure box, in inches, for this slide's layout."""
    top = CONTENT_TOP_SUB if spec.get("subtitle") else CONTENT_TOP_PLAIN
    bottom = CONTENT_BOTTOM - (CALLOUT_H + 0.18 if spec.get("callout") else 0)
    h = bottom - top
    w = RIGHT - LEFT
    layout = spec.get("layout")
    if not layout:
        if spec.get("video"):
            layout = "figure-right" if spec.get("bullets") else "figure-full"
        elif spec.get("figure"):
            layout = "figure-bottom"
        elif spec.get("stage"):
            layout = "figure-right"   # the pipeline strip lives where a figure would
        else:
            layout = "text-only"

    span = (top, bottom)
    if layout in ("assertion-evidence", "ae"):
        # Alley's structure: a sentence headline, then visual evidence, no bullets
        return {"layout": "assertion-evidence", "span": (ASSERT_TOP, bottom),
                "body": None,
                "fig": (BLEED_L, ASSERT_TOP, BLEED_R - BLEED_L, bottom - ASSERT_TOP)}
    if layout == "text-only":
        return {"layout": layout, "span": span, "body": (LEFT, top, w, h), "fig": None}
    if layout == "figure-right":
        bw = w * 0.50
        return {"layout": layout, "span": span,
                "body": (LEFT, top, bw, h),
                "fig": (LEFT + bw + GUTTER, top, BLEED_R - (LEFT + bw + GUTTER), h)}
    if layout == "figure-left":
        fw = w * 0.48
        return {"layout": layout, "span": span,
                "body": (BLEED_L + fw + GUTTER, top, RIGHT - (BLEED_L + fw + GUTTER), h),
                "fig": (BLEED_L, top, fw, h)}
    if layout == "figure-full":
        return {"layout": layout, "span": span, "body": None,
                "fig": (BLEED_L, top, BLEED_R - BLEED_L, h)}
    if layout == "figure-bottom":
        th = min(1.85, h * 0.38)
        return {"layout": layout, "span": span,
                "body": (LEFT, top, w, th),
                "fig": (BLEED_L, top + th + 0.12, BLEED_R - BLEED_L, h - th - 0.12)}
    raise SystemExit(f"unknown layout {layout!r}")


# ---------------------------------------------------------------- helpers

def _is_page_number(shape) -> bool:
    return shape.has_text_frame and Emu(shape.top).inches > 6.6 and Emu(shape.width).inches < 4.5


def _text_shapes(slide):
    return [s for s in slide.shapes if s.has_text_frame and not _is_page_number(s)]


def find_title(slide):
    cands = [s for s in _text_shapes(slide) if Emu(s.top).inches < 1.9]
    return max(cands, key=lambda s: s.width * s.height) if cands else None


MIN_BODY_AREA = 2.0     # in^2; below this it is a label, not the body placeholder


def find_body(slide, title):
    """The bullet region -- never a small fixed label such as the Todolist header.

    Without the area floor, a conclusion slide (whose only other text frame is
    the template's `Todolist & Suggestion from Prof.` caption) hands that label
    back as the body, and the caller then clears it and stretches it across the
    content area.
    """
    cands = [s for s in _text_shapes(slide)
             if s is not title and Emu(s.top).inches >= 1.5
             and Emu(s.width).inches * Emu(s.height).inches >= MIN_BODY_AREA]
    return max(cands, key=lambda s: s.width * s.height) if cands else None


def strip_style(shape) -> None:
    """Drop the theme's <p:style>; it re-applies a text shadow we never asked for."""
    sp = shape._element
    for st in sp.findall('{http://schemas.openxmlformats.org/presentationml/2006/main}style'):
        sp.remove(st)


def place(shape, box) -> None:
    x, y, w, h = box
    shape.left, shape.top, shape.width, shape.height = (
        Inches(x), Inches(y), Inches(w), Inches(h))


def textbox(slide, box, text, size, *, color=None, bold=False, italic=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    x, y, w, h = box
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return tb


def fit(img_box, iw, ih, *, vcenter=True):
    """Aspect-preserving fit. Side layouts centre; stacked layouts hug the top."""
    x, y, w, h = img_box
    scale = min(w / iw, h / ih)
    fw, fh = iw * scale, ih * scale
    dy = (h - fh) / 2 if vcenter else 0.0
    return (x + (w - fw) / 2, y + dy, fw, fh)


# ---------------------------------------------------------------- pieces

def add_subtitle(slide, spec) -> None:
    title = find_title(slide)
    if title is not None:
        place(title, TITLE)
        title.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    textbox(slide, SUBTITLE, spec["subtitle"], 18, color=RED)


def add_assertion(slide, spec) -> None:
    """The claim, at reading size, in place of the small red subtitle."""
    title = find_title(slide)
    if title is not None:
        place(title, TITLE)
        title.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    textbox(slide, ASSERT, spec["subtitle"], 24,
            color=RGBColor(0x1A, 0x1A, 0x1A), bold=True, anchor=MSO_ANCHOR.TOP)


def add_figure(slide, spec, box, *, vcenter=True) -> tuple[float, float, float, float] | None:
    from PIL import Image

    fig = spec["figure"]
    src = Path(fig["src"] if isinstance(fig, dict) else fig)
    if not src.exists():
        raise SystemExit(f"figure not found: {src}")
    fig = fig if isinstance(fig, dict) else {"src": str(src)}

    x, y, w, h = box
    if fig.get("caption"):
        h -= CAPTION_H
    iw, ih = Image.open(src).size
    rect = fit((x, y, w, h), iw, ih, vcenter=vcenter)
    slide.shapes.add_picture(str(src), Inches(rect[0]), Inches(rect[1]),
                             Inches(rect[2]), Inches(rect[3]))
    if fig.get("caption"):
        textbox(slide, (x, rect[1] + rect[3] + 0.06, w, CAPTION_H),
                fig["caption"], 11, color=GREY, italic=True, align=PP_ALIGN.CENTER)
    return rect


BADGE_BG = RGBColor(0x1A, 0x1A, 0x1A)   # the play badge on a video poster
BADGE_H = 0.30


def _play_badge(slide, rect, label: str) -> None:
    """A small marker so nobody mistakes the poster for a still figure."""
    x, y, w, h = rect
    bw = 0.36 + 0.11 * len(label)
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x + 0.10), Inches(y + h - BADGE_H - 0.10),
                                 Inches(bw), Inches(BADGE_H))
    strip_style(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = BADGE_BG
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold = Pt(11), True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _set_playback(movie, opts: dict) -> None:
    """Rewrite the `p:timing` node python-pptx wrote, for autoplay and loop.

    `add_movie` always emits `<p:cond delay="indefinite"/>` - play on click.
    Flipping that one attribute to `delay="0"` is what PowerPoint itself writes
    for "Start: Automatically", and `repeatCount="indefinite"` on the same time
    node is "Loop until Stopped". Both are plain ISO-29500 attributes, so the
    result still passes `office/validate.py`.
    """
    from pptx.oxml.ns import qn

    pic = movie._element
    spid = pic.find(qn("p:nvPicPr")).find(qn("p:cNvPr")).get("id")
    sld = pic.getroottree().getroot()
    tgt = sld.xpath('.//p:timing//p:video//p:spTgt[@spid="%s"]' % spid)
    if not tgt:
        raise SystemExit(
            "python-pptx did not write a p:timing node for this movie. That "
            "happens when the slide already has one wrapped in "
            "<mc:AlternateContent> (a non-OOXML animation, python-pptx issue "
            "#954) - remove the animation from that slide and rebuild.")
    media = tgt[0].getparent().getparent()          # spTgt -> tgtEl -> cMediaNode
    cTn = media.find(qn("p:cTn"))
    if opts.get("autoplay"):
        cTn.find(qn("p:stCondLst")).find(qn("p:cond")).set("delay", "0")
    if opts.get("loop"):
        cTn.set("repeatCount", "indefinite")
    if opts.get("mute"):
        media.set("mute", "1")
    if opts.get("volume") is not None:
        pct = max(0.0, min(100.0, float(opts["volume"])))
        media.set("vol", str(int(round(pct * 1000))))   # 1000ths of a percent
    if opts.get("fullscreen"):
        media.getparent().set("fullScrn", "1")


def add_video(slide, spec, box, *, vcenter=True) -> tuple[float, float, float, float]:
    """Embed a movie the way PowerPoint expects to find one.

    Three things this gets right that a bare `add_movie` call does not: the
    media part is typed by extension instead of `video/unknown`, the file is
    refused up front when PowerPoint cannot decode it, and a real frame is
    pulled from the video as the poster so the slide is not a grey speaker icon.
    """
    import video as video_mod
    from PIL import Image

    v = spec["video"]
    v = {"src": v} if isinstance(v, str) else dict(v)
    src = Path(v["src"])
    if not src.exists():
        raise SystemExit(f"video not found: {src}")

    info = video_mod.probe(src)
    blocking = [m for lv, m in video_mod.problems(src, info) if lv == "error"]
    if blocking:
        raise SystemExit(
            f"{src} will not play in PowerPoint:\n  - "
            + "\n  - ".join(blocking)
            + f"\nFix it once with:\n  {video_mod.fix_command(src)}")

    still = Path(v["poster"]) if v.get("poster") else \
        video_mod.auto_poster(src, v.get("poster_at"))
    if not still.exists():
        raise SystemExit(f"poster image not found: {still}")

    x, y, w, h = box
    if v.get("caption"):
        h -= CAPTION_H
    iw, ih = Image.open(still).size
    rect = fit((x, y, w, h), iw, ih, vcenter=vcenter)
    movie = slide.shapes.add_movie(
        str(src), Inches(rect[0]), Inches(rect[1]), Inches(rect[2]), Inches(rect[3]),
        poster_frame_image=str(still), mime_type=video_mod.mime_for(src))
    _set_playback(movie, v)

    if v.get("badge", True):
        secs = int(info.get("duration") or 0)
        _play_badge(slide, rect,
                    f"\u25b6 {secs // 60}:{secs % 60:02d}" if secs else "\u25b6")
    if v.get("caption"):
        textbox(slide, (x, rect[1] + rect[3] + 0.06, w, CAPTION_H),
                v["caption"], 11, color=GREY, italic=True, align=PP_ALIGN.CENTER)
    return rect


def add_annotations(slide, spec, rect) -> None:
    """Coordinates are fractions of the placed picture, so they survive resizing."""
    fx, fy, fw, fh = rect
    for a in spec.get("annotations", []):
        kind = a.get("type", "box")
        colour = RGBColor.from_string(a.get("color", "C00000"))
        at = a["at"]
        if kind in ("box", "circle"):
            x, y, w, h = at
            shp = slide.shapes.add_shape(
                MSO_SHAPE.OVAL if kind == "circle" else MSO_SHAPE.RECTANGLE,
                Inches(fx + x * fw), Inches(fy + y * fh),
                Inches(w * fw), Inches(h * fh))
            strip_style(shp)
            shp.fill.background()
            shp.line.color.rgb = colour
            shp.line.width = Pt(2.25)
            if a.get("dash"):
                shp.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
            shp.shadow.inherit = False
            if a.get("text"):
                shp.text_frame.text = a["text"]
                r = shp.text_frame.paragraphs[0].runs[0]
                r.font.size, r.font.bold, r.font.color.rgb = Pt(12), True, colour
        elif kind == "arrow":
            x0, y0, x1, y1 = at
            conn = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(fx + x0 * fw), Inches(fy + y0 * fh),
                Inches(fx + x1 * fw), Inches(fy + y1 * fh))
            conn.line.color.rgb = colour
            conn.line.width = Pt(2.25)
            _arrowhead(conn)
        elif kind == "label":
            x, y = at[0], at[1]
            textbox(slide, (fx + x * fw, fy + y * fh, 3.4, 0.36),
                    a.get("text", ""), a.get("size", 13), color=colour,
                    bold=True, wrap=False)


def _arrowhead(connector) -> None:
    from pptx.oxml.ns import qn
    ln = connector.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)


def add_callout(slide, spec) -> None:
    c = spec["callout"]
    label, text = (c.get("label"), c["text"]) if isinstance(c, dict) else (None, c)
    y = CONTENT_BOTTOM - CALLOUT_H
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(BLEED_L), Inches(y),
                                 Inches(BLEED_R - BLEED_L), Inches(CALLOUT_H))
    strip_style(shp)
    shp.fill.background()
    shp.line.color.rgb = RED
    shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.16)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    if label:
        r = p.add_run()
        r.text = f"{label}:  "
        r.font.size, r.font.bold, r.font.color.rgb = Pt(15), True, RED
    r = p.add_run()
    r.text = text
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def add_matrix(slide, spec, box) -> None:
    m = spec["matrix"]
    header, rows = m["header"], m["rows"]
    hi = m.get("highlight_row")
    x, y, w, h = box
    n_rows, n_cols = len(rows) + 1, len(header)
    height = min(h, 0.42 + 0.40 * len(rows))
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                Inches(w), Inches(height))
    table = gf.table
    table.first_row = True
    for c, txt in enumerate(header):
        cell = table.cell(0, c)
        cell.text = str(txt)
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEAD_BG
        _style_cell(cell, 13, bold=True, color=HEAD_FG,
                    align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    for r, row in enumerate(rows, start=1):
        for c, txt in enumerate(row[:n_cols]):
            cell = table.cell(r, c)
            cell.text = str(txt)
            if hi is not None and r == hi:
                cell.fill.solid()
                cell.fill.fore_color.rgb = ROW_HI
            else:
                cell.fill.background()
            _style_cell(cell, 12.5, bold=(hi is not None and r == hi),
                        align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)


def _style_cell(cell, size, *, bold=False, color=None, align=PP_ALIGN.LEFT):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.08)
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            if color is not None:
                r.font.color.rgb = color


def equation_height(eq: dict) -> float:
    """Vertical band an equation needs, including its where-list and captions."""
    h = 1.15
    h += 0.24 * len(eq.get("where") or [])
    h += 0.28 * len(eq.get("captions") or {})
    if eq.get("label"):
        h += 0.26
    return min(h, 3.4)


def _equation_image(eq: dict, outline_dir: Path, tag: str) -> Path:
    """A cropped equation is used as given; a rebuilt one is rendered now."""
    if eq.get("src"):
        src = Path(eq["src"])
        if not src.exists():
            raise SystemExit(f"equation image not found: {src}")
        return src
    parts = eq.get("parts")
    if not parts:
        raise SystemExit("equation needs either `src` (quote it) or `parts` (explain it)")
    import equation as equation_mod
    pairs = [(p[0], (p[1] if len(p) > 1 else "111111").lstrip("#").upper())
             for p in parts]
    # scratch, like the render PNGs: the deck is the only thing that lands in
    # the user's directory
    import tempfile
    out = Path(tempfile.gettempdir()) / "acm-eq" / f"{outline_dir.name}-{tag}.png"
    out.parent.mkdir(parents=True, exist_ok=True)
    return equation_mod.render(pairs, out)


def add_equation(slide, spec, box, outline_dir: Path, tag: str) -> None:
    eq = spec["equation"]
    if isinstance(eq, str):                      # tolerate the old plain-text form
        eq = {"parts": [[eq]]}
    from PIL import Image

    x, y, w, h = box
    where = eq.get("where") or []
    captions = eq.get("captions") or {}
    below = 0.24 * len(where) + 0.28 * len(captions) + (0.26 if eq.get("label") else 0)

    img = _equation_image(eq, outline_dir, tag)
    iw, ih = Image.open(img).size
    rect = fit((x, y, w, max(0.5, h - below)), iw, ih, vcenter=True)
    slide.shapes.add_picture(str(img), Inches(rect[0]), Inches(rect[1]),
                             Inches(rect[2]), Inches(rect[3]))
    if eq.get("annotations"):
        add_annotations(slide, {"annotations": eq["annotations"]}, rect)

    cy = rect[1] + rect[3] + 0.05
    if eq.get("label"):
        textbox(slide, (x, cy, w, 0.24), eq["label"], 11,
                color=GREY, italic=True, align=PP_ALIGN.CENTER)
        cy += 0.26
    for colour, text in captions.items():
        textbox(slide, (x, cy, w, 0.26), text, 12,
                color=RGBColor.from_string(colour.lstrip("#").upper()),
                bold=True, align=PP_ALIGN.CENTER)
        cy += 0.28
    for line in where:
        textbox(slide, (x + 0.10, cy, w - 0.10, 0.22), line, 12, color=GREY)
        cy += 0.24


# ---------------------------------------------------------------- recurring strip

def add_stage(slide, spec, stage_figure, box) -> None:
    """The method-overview figure, repeated, with a box on the stage being told.

    Both reference decks do this on every method slide; naming the figure once
    at the top of the outline keeps it out of every slide entry.
    """
    from PIL import Image

    src = Path(stage_figure["src"])
    if not src.exists():
        raise SystemExit(f"stage_figure not found: {src}")
    name = spec["stage"]
    at = (stage_figure.get("stages") or {}).get(name)
    if at is None:
        raise SystemExit(f"stage {name!r} is not in stage_figure.stages")

    x, y, w, _ = box
    iw, ih = Image.open(src).size
    rect = fit((x, y, w, STAGE_H), iw, ih, vcenter=False)
    slide.shapes.add_picture(str(src), Inches(rect[0]), Inches(rect[1]),
                             Inches(rect[2]), Inches(rect[3]))
    add_annotations(slide, {"annotations": [
        {"type": "box", "at": at, "color": stage_figure.get("color", "C00000")}]}, rect)
    if stage_figure.get("caption"):
        textbox(slide, (x, rect[1] + rect[3] + 0.04, w, 0.24),
                stage_figure["caption"], 10, color=GREY, italic=True,
                align=PP_ALIGN.CENTER)


def add_divider(slide, reg) -> None:
    """The thin rule between a figure column and its text, as in the lab decks."""
    if reg["fig"] is None or reg["body"] is None:
        return
    bx, bw = reg["body"][0], reg["body"][2]
    fx = reg["fig"][0]
    x = (bx + bw + fx) / 2 if fx > bx else (fx + reg["fig"][2] + bx) / 2
    top, bottom = reg["span"]
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x), Inches(top + 0.05),
                                      Inches(x), Inches(bottom - 0.05))
    line.line.color.rgb = DIVIDER
    line.line.width = Pt(1.0)


# ---------------------------------------------------------------- driver

def flatten(outline: dict) -> list[dict]:
    flat = []
    if outline.get("meta"):
        flat.append({"role": "cover"})
    if outline.get("summary"):
        flat.append({"role": "summary"})
    flat.extend(outline.get("slides", []))
    return flat


VISUAL_KEYS = ("subtitle", "figure", "video", "annotations", "callout",
               "matrix", "equation", "stage", "divider")


STAMP = "acm-composed"          # so a second pass cannot silently double every shape


def compose(outline_path: Path, deck: Path, out: Path | None = None,
            force: bool = False) -> Path:
    outline = json.loads(outline_path.read_text(encoding="utf-8"))
    flat = flatten(outline)
    prs = Presentation(str(deck))
    if prs.core_properties.content_status == STAMP and not force:
        raise SystemExit(
            f"{deck} has already been composed - running this again stacks a "
            f"second figure, caption and callout on top of the first. Rebuild "
            f"with build_from_outline.py, then compose. (--force overrides.)")
    if len(prs.slides._sldIdLst) != len(flat):
        raise SystemExit(
            f"{deck} has {len(prs.slides._sldIdLst)} slides but the outline "
            f"describes {len(flat)} - rebuild with build_from_outline.py first")

    touched = 0
    stage_figure = outline.get("stage_figure")
    outline_dir = outline_path.resolve().parent

    for idx, (slide, spec) in enumerate(zip(prs.slides, flat), start=1):
        if not any(spec.get(k) for k in VISUAL_KEYS):
            continue
        touched += 1
        reg = regions(spec)
        if spec.get("subtitle"):
            if reg["layout"] == "assertion-evidence":
                add_assertion(slide, spec)
            else:
                add_subtitle(slide, spec)

        # an equation shares the body column with the bullets, so claim its band
        # before the body placeholder is resized -- that is what stops overlap
        eq_box = None
        if spec.get("equation") and reg["body"]:
            bx, by, bw, bh = reg["body"]
            eq = spec["equation"]
            band = equation_height(eq if isinstance(eq, dict) else {})
            if spec.get("bullets"):
                band = min(band, bh - 0.9)
                eq_box = (bx, by + bh - band, bw, band)
                reg["body"] = (bx, by, bw, bh - band - 0.15)
            else:
                eq_box = (bx, by, bw, bh)
                reg["body"] = (bx, by, 0.4, 0.3)

        # a table slide's content is the table build_from_outline already filled;
        # there is no bullet region to move, clear or shrink
        body = None if spec.get("table") else find_body(slide, find_title(slide))
        if body is not None and not spec.get("bullets"):
            body.text_frame.clear()
        if body is not None:
            if reg["body"] is None:
                body.text_frame.clear()
                place(body, (LEFT, CONTENT_TOP_PLAIN, 0.4, 0.3))
            else:
                place(body, reg["body"])
                if reg["layout"] in ("figure-right", "figure-left"):
                    for para in body.text_frame.paragraphs:
                        para.alignment = PP_ALIGN.LEFT

        if eq_box:
            add_equation(slide, spec, eq_box, outline_dir, str(idx))
        if spec.get("matrix"):
            add_matrix(slide, spec, reg["body"] or reg["fig"])
        if spec.get("stage") and reg["fig"]:
            if not stage_figure:
                raise SystemExit(f"slide {idx} sets `stage` but the outline has no "
                                 f"top-level `stage_figure`")
            add_stage(slide, spec, stage_figure, reg["fig"])
        elif spec.get("video") and reg["fig"]:
            if spec.get("figure"):
                raise SystemExit(f"slide {idx} has both `figure` and `video` - "
                                 f"one exhibit per slide; split them")
            rect = add_video(slide, spec, reg["fig"],
                             vcenter=reg["layout"] in ("figure-right", "figure-left",
                                                       "assertion-evidence"))
            if spec.get("annotations"):
                add_annotations(slide, spec, rect)
        elif spec.get("figure") and reg["fig"]:
            rect = add_figure(slide, spec, reg["fig"],
                              vcenter=reg["layout"] in ("figure-right", "figure-left",
                                                        "assertion-evidence"))
            if spec.get("annotations"):
                add_annotations(slide, spec, rect)
        if spec.get("divider"):
            add_divider(slide, reg)
        if spec.get("callout"):
            add_callout(slide, spec)

    out = out or deck
    prs.core_properties.content_status = STAMP
    prs.save(str(out))
    print(f"composed {touched}/{len(flat)} slides -> {out}")
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("outline")
    ap.add_argument("deck")
    ap.add_argument("-o", "--output", help="default: edit the deck in place")
    ap.add_argument("--force", action="store_true",
                    help="compose a deck that was already composed once")
    a = ap.parse_args()
    compose(Path(a.outline), Path(a.deck),
            Path(a.output) if a.output else None, force=a.force)


if __name__ == "__main__":
    main()
