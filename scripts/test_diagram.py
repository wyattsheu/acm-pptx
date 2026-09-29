#!/usr/bin/env python3
"""Smoke-test native diagram layout and editability in a real PPTX package."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parent.parent


def run(*args: str, cwd: Path) -> None:
    subprocess.run([sys.executable, *args], cwd=cwd, check=True,
                   stdout=subprocess.DEVNULL)


def main() -> None:
    outline = {
        "slides": [
            {
                "role": "others",
                "title": "Editable left-to-right pipeline",
                "diagram": {
                    "direction": "LR",
                    "nodes": [
                        {"id": "in", "text": "Input", "kind": "terminator"},
                        {"id": "work", "text": "Transform", "kind": "process"},
                        {"id": "ok", "text": "Valid?", "kind": "decision"},
                        {"id": "out", "text": "Output", "kind": "terminator",
                         "accent": True},
                    ],
                    "edges": [
                        {"id": "read", "from": "in", "to": "work"},
                        {"id": "check", "from": "work", "to": "ok"},
                        {"id": "emit", "from": "ok", "to": "out", "label": "yes"},
                    ],
                    "groups": [
                        {"id": "core", "label": "Core", "nodes": ["work", "ok"]}
                    ],
                },
                "notes": "這是一個原生可編輯流程圖。節點與箭頭都可單獨修改。移動節點時連接線端點會跟著移動。",
            },
            {
                "role": "others",
                "title": "Editable top-to-bottom branch",
                "diagram": {
                    "direction": "TB",
                    "nodes": [
                        {"id": "start", "text": "Start", "kind": "terminator"},
                        {"id": "route", "text": "Route", "kind": "decision"},
                        {"id": "fast", "text": "Fast path", "kind": "process"},
                        {"id": "safe", "text": "Safe path", "kind": "process"},
                    ],
                    "edges": [
                        {"from": "start", "to": "route"},
                        {"from": "route", "to": "fast", "label": "ready"},
                        {"from": "route", "to": "safe", "label": "fallback",
                         "dashed": True},
                    ],
                },
                "notes": "這頁測試上下方向與分支。決策節點使用菱形。備援路徑用虛線表達，不靠顏色區分。",
            },
        ]
    }
    expected = [
        (4, ["read", "check", "emit"], ["core"]),
        (4, ["e1", "e2", "e3"], []),
    ]

    with tempfile.TemporaryDirectory(prefix="acm-diagram-test-") as tmp:
        work = Path(tmp)
        outline_path = work / "outline.json"
        deck = work / "diagram.pptx"
        outline_path.write_text(json.dumps(outline, ensure_ascii=False), encoding="utf-8")
        run(str(ROOT / "scripts" / "build_from_outline.py"), str(outline_path),
            "-o", str(deck), cwd=work)
        run(str(ROOT / "scripts" / "compose.py"), str(outline_path), str(deck), cwd=work)
        run(str(ROOT / "scripts" / "qa_check.py"), str(outline_path), str(deck), cwd=work)

        prs = Presentation(str(deck))
        assert len(prs.slides) == 2
        native_nodes = attached_edges = 0
        for index, (slide, (node_count, edge_ids, group_ids)) in enumerate(
                zip(prs.slides, expected), start=1):
            by_name = {shape.name: shape for shape in slide.shapes}
            nodes = [name for name in by_name if name.startswith("diagram:node:")]
            assert len(nodes) == node_count, (index, nodes)
            native_nodes += len(nodes)
            for edge_id in edge_ids:
                edge = by_name[f"diagram:edge:{edge_id}"]
                assert edge._element.xpath(".//a:stCxn")
                assert edge._element.xpath(".//a:endCxn")
                attached_edges += 1
            for group_id in group_ids:
                assert f"diagram:group:{group_id}" in by_name

        with zipfile.ZipFile(deck) as archive:
            for index in (1, 2):
                xml = archive.read(f"ppt/slides/slide{index}.xml").decode("utf-8")
                assert "<p:pic>" not in xml
        print(f"editable diagram smoke test PASSED "
              f"({native_nodes} nodes, {attached_edges} attached edges, 0 pictures)")


if __name__ == "__main__":
    main()
