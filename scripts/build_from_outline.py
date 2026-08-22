"""Build an ACM Lab progress-report deck from an outline JSON.

The lab template (assets/acm_template.pptx) is 22 slides, each with its own
layout carrying the section colour bar and breadcrumb. This script never
invents a layout: it clones the template slide whose *role* matches each
outline entry, so the breadcrumb, colour, and typography stay exactly as the
lab designed them, then replaces only the text.

Usage:
    python build_from_outline.py outline.json -o report.pptx
    python build_from_outline.py outline.json -o report.pptx --template path/to/tpl.pptx
    python build_from_outline.py --roles          # print the role table and exit

Outline schema: see references/outline-schema.md
"""

from __future__ import annotations

import argparse
import contextlib
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from add_slide import duplicate_slide  # noqa: E402
from office.helpers import rezip, safe_extract  # noqa: E402

DEFAULT_TEMPLATE = SCRIPT_DIR.parent / "assets" / "acm_template.pptx"

# role -> (template slide number, human description)
ROLES: dict[str, tuple[int, str]] = {
    "cover":               (1,  "Cover — Progress Report / date / advisor / student"),
    "summary":             (2,  "Summary — 4-row status table + numbered weekly task table"),
    "project_summary":     (3,  "Project · agenda page (Goal/Architecture/Status/Discussion)"),
    "project_why":         (4,  "Project · Why This Project Matters"),
    "project_prev":        (11, "Previous Method (blue Research bar — reuse if needed)"),
    "project_pipeline":    (5,  "Project · Current Pipeline Overview"),
    "project_results":     (6,  "Project · Progress & Results"),
    "project_insight":     (7,  "Project · Insight / discussion items"),
    "project_conclusion":  (8,  "Project · Conclusion (numbered todolist table)"),
    "research_summary":    (9,  "Research · agenda page"),
    "research_why":        (10, "Research · Why This Research Matters"),
    "research_prev":       (11, "Research · Previous Method"),
    "research_pipeline":   (12, "Research · Current Pipeline Overview"),
    "research_results":    (13, "Research · Progress & Results"),
    "research_insight":    (14, "Research · Insight / discussion items"),
    "research_conclusion": (15, "Research · Conclusion (numbered todolist table)"),
    "paper_list":          (16, "Paper Study · reading list table (done / no / venue / title)"),
    "paper_intro":         (17, "Paper · Introduction"),
    "paper_related":       (18, "Paper · Related work"),
    "paper_method":        (19, "Paper · Proposed Method"),
    "paper_results":       (20, "Paper · Experimental Results"),
    "paper_conclusion":    (21, "Paper · Conclusion"),
    "others":              (22, "Others"),
}

PAGE_NUM_RE = re.compile(r"^[‹<]#[›>]$")


# --------------------------------------------------------------------------
# deck assembly (package level)
# --------------------------------------------------------------------------

def assemble(template: Path, roles: list[str], out: Path) -> None:
    """Clone the template slide for each role, in order, into a fresh deck."""
    tmp = Path(tempfile.mkdtemp())
    try:
        unpacked = tmp / "unpacked"
        import zipfile
        with zipfile.ZipFile(template) as zf:
            safe_extract(zf, unpacked)

        created = []
        for role in roles:
            src = f"slide{ROLES[role][0]}.xml"
            with contextlib.redirect_stdout(io.StringIO()):
                dest = duplicate_slide(unpacked, src)
            created.append(Path(dest).name)

        _rewrite_sld_id_lst(unpacked, created)
        subprocess.run(
            [sys.executable, str(SCRIPT_DIR / "clean.py"), str(unpacked)],
            check=True, capture_output=True,
        )
        rezip(unpacked, out)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _rewrite_sld_id_lst(unpacked: Path, slide_files: list[str]) -> None:
    """Keep only the freshly cloned slides, in the given order."""
    rels_path = unpacked / "ppt" / "_rels" / "presentation.xml.rels"
    rels = rels_path.read_text(encoding="utf-8")
    rid_of = {}
    for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="slides/([^"]+)"', rels):
        rid_of[m.group(2)] = m.group(1)
    for m in re.finditer(r'<Relationship[^>]*Target="slides/([^"]+)"[^>]*Id="([^"]+)"', rels):
        rid_of.setdefault(m.group(1), m.group(2))

    entries = []
    for i, name in enumerate(slide_files):
        rid = rid_of.get(name)
        if rid is None:
            raise SystemExit(f"no relationship found for {name}")
        entries.append(f'<p:sldId id="{900 + i}" r:id="{rid}"/>')

    pres_path = unpacked / "ppt" / "presentation.xml"
    xml = pres_path.read_text(encoding="utf-8")
    xml = re.sub(
        r"<p:sldIdLst>.*?</p:sldIdLst>",
        "<p:sldIdLst>" + "".join(entries) + "</p:sldIdLst>",
        xml,
        count=1,
        flags=re.DOTALL,
    )
    pres_path.write_text(xml, encoding="utf-8")


# --------------------------------------------------------------------------
# text filling (python-pptx level)
# --------------------------------------------------------------------------

def _is_page_number(shape) -> bool:
    return shape.has_text_frame and PAGE_NUM_RE.match(shape.text_frame.text.strip() or "x") is not None


def _text_shapes(slide):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and not _is_page_number(sh):
            out.append(sh)
    return out


def _find_title(slide):
    try:
        if slide.shapes.title is not None:
            return slide.shapes.title
    except (AttributeError, ValueError):
        pass
    cands = [sh for sh in _text_shapes(slide) if sh.is_placeholder]
    if cands:
        return min(cands, key=lambda s: (s.top or 0))
    return None


def _find_body(slide, title_shape):
    """Biggest non-title text frame — where the bullets live."""
    cands = [
        sh for sh in _text_shapes(slide)
        if sh is not title_shape and (sh.height or 0) * (sh.width or 0) > 0
    ]
    if not cands:
        return None
    return max(cands, key=lambda s: (s.height or 0) * (s.width or 0))


def set_paragraphs(tf, items) -> None:
    """Replace a text frame's paragraphs, reusing existing ones as style prototypes.

    items: list of str, or of {"text": str, "level": int, "bold": bool}
    """
    norm = []
    for it in items:
        if isinstance(it, str):
            norm.append({"text": it, "level": 0})
        else:
            norm.append({"text": it.get("text", ""), "level": int(it.get("level", 0)),
                         "bold": it.get("bold")})

    # This template expresses hierarchy with marL/indent/buChar at lvl="0",
    # not with the lvl attribute — so derive tiers from the left margin and
    # keep each tier's paragraph as a style prototype.
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def _marl(p_el) -> int:
        pPr = p_el.find(f"{A}pPr")
        if pPr is None:
            return 0
        try:
            return int(pPr.get("marL") or 0)
        except ValueError:
            return 0

    by_marl: dict[int, object] = {}
    for p in tf.paragraphs:
        by_marl.setdefault(_marl(p._p), deepcopy(p._p))
    if not by_marl:
        by_marl[0] = deepcopy(tf.paragraphs[0]._p)
    tiers = [by_marl[k] for k in sorted(by_marl)]
    max_tier = len(tiers) - 1

    body = tf._txBody
    for p in list(tf.paragraphs):
        body.remove(p._p)

    for spec in norm:
        tier = max(0, min(spec["level"], max_tier))
        body.append(deepcopy(tiers[tier]))

    from pptx.text.text import _Paragraph
    for spec, p_el in zip(norm, list(body.findall(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}p"))):
        para = _Paragraph(p_el, tf)
        runs = para.runs
        if not runs:
            para.text = spec["text"]
        else:
            runs[0].text = spec["text"]
            for extra in runs[1:]:
                extra._r.getparent().remove(extra._r)
        if spec.get("bold") is not None and para.runs:
            para.runs[0].font.bold = bool(spec["bold"])


def set_cell(cell, text: str) -> None:
    """Write a table cell, preserving the run formatting already there."""
    tf = cell.text_frame
    para = tf.paragraphs[0]
    for extra in list(tf.paragraphs)[1:]:
        tf._txBody.remove(extra._p)
    runs = para.runs
    if not runs:
        para.text = text
    else:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)


def _tables(slide):
    return [sh.table for sh in slide.shapes if sh.has_table]


def fill_cover(slide, meta: dict) -> None:
    for sh in _text_shapes(slide):
        txt = sh.text_frame.text
        if "Progress Report" in txt:
            set_paragraphs(sh.text_frame, [
                meta.get("title", "Progress Report"),
                meta.get("date", ""),
            ])
        elif "Advisor" in txt or "Student" in txt:
            set_paragraphs(sh.text_frame, [
                f"Advisor: {meta.get('advisor', '')}",
                f"Student: {meta.get('student', '')}",
            ])


def fill_summary(slide, spec: dict) -> None:
    title = _find_title(slide)
    if title is not None and spec.get("title"):
        set_paragraphs(title.text_frame, [spec["title"]])

    tabs = _tables(slide)
    status = spec.get("rows") or []
    tasks = spec.get("tasks") or []

    # the 4-row table is the status block; the 8-row one is the task list
    status_tab = next((t for t in tabs if len(t.rows) == len(status)), None)
    if status_tab is None:
        status_tab = min(tabs, key=lambda t: len(t.rows)) if tabs else None
    task_tab = next((t for t in tabs if t is not status_tab), None)

    if status_tab is not None:
        for r, row in enumerate(status):
            if r >= len(status_tab.rows):
                break
            for c, val in enumerate(row[: len(status_tab.columns)]):
                set_cell(status_tab.cell(r, c), str(val))

    if task_tab is not None:
        _fill_numbered(task_tab, tasks)

    if spec.get("note"):
        for sh in _text_shapes(slide):
            if "summarize" in sh.text_frame.text.lower():
                set_paragraphs(sh.text_frame, [spec["note"]])


def _fill_numbered(table, items) -> None:
    """Numbered 2-column table: col0 = index, col1 = text. Blank out unused rows."""
    n_text_col = len(table.columns) - 1
    for r in range(len(table.rows)):
        val = items[r] if r < len(items) else ""
        set_cell(table.cell(r, n_text_col), str(val))


def fill_table_slide(slide, spec: dict) -> None:
    title = _find_title(slide)
    if title is not None and spec.get("title"):
        set_paragraphs(title.text_frame, [spec["title"]])
    tabs = _tables(slide)
    if not tabs:
        return
    table = tabs[0]
    rows = spec.get("table")
    if rows and isinstance(rows[0], (list, tuple)):
        for r, row in enumerate(rows[: len(table.rows)]):
            for c, val in enumerate(row[: len(table.columns)]):
                set_cell(table.cell(r, c), str(val))
        for r in range(len(rows), len(table.rows)):
            for c in range(len(table.columns)):
                set_cell(table.cell(r, c), "")
    elif rows:
        _fill_numbered(table, rows)


def fill_generic(slide, spec: dict) -> None:
    title = _find_title(slide)
    if title is not None:
        set_paragraphs(title.text_frame, [spec.get("title", title.text_frame.text)])
    body = _find_body(slide, title)
    if body is not None and spec.get("bullets"):
        set_paragraphs(body.text_frame, spec["bullets"])
    elif body is not None and spec.get("bullets") == []:
        set_paragraphs(body.text_frame, [""])
    # drop the grey "if you don't have project..." hint once real content lands
    for sh in _text_shapes(slide):
        t = sh.text_frame.text.lower()
        if "just show this slides" in t and spec.get("bullets"):
            set_paragraphs(sh.text_frame, [spec.get("subtitle", "")])


def fill(out: Path, outline: dict) -> None:
    from pptx import Presentation

    prs = Presentation(str(out))
    for slide, spec in zip(prs.slides, outline["_flat"]):
        role = spec["role"]
        if role == "cover":
            fill_cover(slide, outline.get("meta", {}))
        elif role == "summary":
            fill_summary(slide, spec)
        elif spec.get("table") is not None:
            fill_table_slide(slide, spec)
        else:
            fill_generic(slide, spec)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(str(out))


# --------------------------------------------------------------------------

def flatten(outline: dict) -> list[dict]:
    flat = []
    if outline.get("meta"):
        flat.append({"role": "cover"})
    if outline.get("summary"):
        s = dict(outline["summary"])
        s["role"] = "summary"
        s.setdefault("title", "Summary")
        flat.append(s)
    for s in outline.get("slides", []):
        if s["role"] not in ROLES:
            raise SystemExit(
                f"unknown role {s['role']!r}. Valid roles:\n  " + "\n  ".join(ROLES)
            )
        flat.append(s)
    if not flat:
        raise SystemExit("outline produced no slides")
    return flat


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("outline", nargs="?", help="outline JSON")
    ap.add_argument("-o", "--output", default="report.pptx")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    ap.add_argument("--roles", action="store_true", help="print the role table and exit")
    args = ap.parse_args()

    if args.roles:
        for k, (n, desc) in ROLES.items():
            print(f"{k:22} -> template slide {n:2}  {desc}")
        return

    if not args.outline:
        ap.error("outline is required (or pass --roles)")

    outline = json.loads(Path(args.outline).read_text(encoding="utf-8"))
    outline["_flat"] = flatten(outline)

    out = Path(args.output)
    assemble(Path(args.template), [s["role"] for s in outline["_flat"]], out)
    fill(out, outline)
    print(f"Wrote {out} ({len(outline['_flat'])} slides)")


if __name__ == "__main__":
    main()
